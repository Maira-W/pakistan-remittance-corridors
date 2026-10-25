"""Turn the two raw files into clean, checked tables (data/interim).

Every check raises an error instead of printing a warning: a number that does not add up must stop the build.
"""

from __future__ import annotations

import json
import re

import openpyxl
import pandas as pd
import pypdfium2 as pdfium

from . import config

UAE_LEAVES = ["Dubai", "Abu Dhabi", "Sharjah", "Others"]
OTHER_GCC_LEAVES = ["Bahrain", "Kuwait", "Qatar", "Oman"]
EU_LEAVES = ["Germany", "France", "Netherland", "Spain", "Italy", "Greece", "Sweden", "Denmark", "Ireland", "Belgium"]


def _month_of(header) -> tuple[pd.Timestamp | None, str | None]:
    """A header cell is a month (a date, or text like 'Jan-24' or 'Mar-24R' for a revised month) or a fiscal year ('FY22')."""
    if hasattr(header, "year"):
        return pd.Timestamp(header.year, header.month, 1), None
    text = str(header).strip()
    if re.fullmatch(r"FY\d\d", text):
        return None, text
    match = re.fullmatch(r"([A-Za-z]{3})-(\d\d)R?", text)
    if not match:
        raise ValueError(f"unrecognised column header {header!r}")
    return pd.to_datetime(f"{match.group(1)}-20{match.group(2)}", format="%b-%Y"), None


def read_workbook() -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Monthly amounts by SBP row (million US dollars), the fiscal-year columns SBP prints, and the checks made."""
    path = config.RAW / config.ARCHIVE_FILE
    ws = openpyxl.load_workbook(path, data_only=True)[config.ARCHIVE_SHEET]
    header = [c.value for c in ws[5]][3:]
    parsed = [_month_of(h) for h in header]
    rows: dict[str, list] = {}
    for r in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=True):
        name = r[2] if r[2] else r[1]
        if isinstance(name, str) and name.strip() == "Total":
            rows["Total"] = list(r[3:3 + len(header)])
        elif isinstance(name, str) and (name in config.LEAF_TO_CORRIDOR or name in ("U.A.E.", "EU Countries", "Other GCC Countries")):
            if name in rows:
                raise ValueError(f"row {name!r} appears twice")
            rows[name] = list(r[3:3 + len(header)])
    missing = [n for n in [*config.LEAF_TO_CORRIDOR, "Total", "U.A.E.", "EU Countries", "Other GCC Countries"] if n not in rows]
    if missing:
        raise ValueError(f"rows missing from the workbook: {missing}")

    months = [m for m, _ in parsed if m is not None]
    month_cols = [i for i, (m, _) in enumerate(parsed) if m is not None]
    fy_cols = {fy: i for i, (_, fy) in enumerate(parsed) if fy is not None}
    wide = pd.DataFrame({name: [vals[i] for i in month_cols] for name, vals in rows.items()}, index=pd.DatetimeIndex(months, name="month"))
    wide = wide.astype(float)

    checks: dict = {}
    expected = pd.date_range(wide.index.min(), wide.index.max(), freq="MS")
    if list(wide.index) != list(expected):
        raise ValueError("months are not continuous")
    if wide.isna().any().any() or (wide < 0).any().any():
        raise ValueError("the workbook has blank or negative amounts")
    checks["months"] = [str(wide.index.min().date()), str(wide.index.max().date()), len(wide)]

    leaves = list(config.LEAF_TO_CORRIDOR)
    diff = (wide[leaves].sum(axis=1) - wide["Total"]).abs().max()
    checks["max_gap_leaves_vs_total"] = round(float(diff), 6)
    if diff > 0.05:
        raise ValueError(f"countries do not add up to the total (largest gap {diff:.3f})")
    for group, parts in (("U.A.E.", UAE_LEAVES), ("Other GCC Countries", OTHER_GCC_LEAVES), ("EU Countries", EU_LEAVES)):
        gap = (wide[parts].sum(axis=1) - wide[group]).abs().max()
        checks[f"max_gap_{group}"] = round(float(gap), 6)
        if gap > 0.05:
            raise ValueError(f"{group} is not the sum of its countries (largest gap {gap:.3f})")

    # SBP prints a fiscal-year total beside each year of months: the months must add up to it
    fy_totals = {}
    for fy, col in fy_cols.items():
        year = 2000 + int(fy[2:])
        window = wide.loc[f"{year - 1}-07-01":f"{year}-06-01"]
        if len(window) != 12:
            continue                                  # FY19 starts before the data
        for name in ("Total", "USA", "Saudi Arabia", "Dubai"):
            printed = float(rows[name][col])
            gap = abs(window[name].sum() - printed)
            if gap > 0.5:
                raise ValueError(f"{fy} {name}: months add to {window[name].sum():.1f} but SBP prints {printed:.1f}")
        fy_totals[fy] = round(float(window["Total"].sum()), 1)
    checks["fiscal_years_confirmed_against_sbp_totals"] = fy_totals

    fy_printed = pd.DataFrame({fy: {name: rows[name][col] for name in rows} for fy, col in fy_cols.items()}).T
    return wide, fy_printed, checks


def corridor_table(wide: pd.DataFrame) -> pd.DataFrame:
    """Long table of months by corridor, plus the region of each corridor."""
    frame = wide[list(config.LEAF_TO_CORRIDOR)].T.groupby(config.LEAF_TO_CORRIDOR).sum().T
    frame = frame[config.CORRIDORS]
    out = frame.reset_index().melt(id_vars="month", var_name="corridor", value_name="usd_million")
    out["region"] = out["corridor"].map(config.CORRIDOR_REGION)
    return out.sort_values(["month", "corridor"]).reset_index(drop=True)


# ---------------------------------------------------------------------------
# The bulletin table: FY25 and FY26 by country
# ---------------------------------------------------------------------------
NUM = r"-?[\d,]+\.\d"
ROW = re.compile(rf"^\s*(?:\d+\.)?\s*([A-Za-z][A-Za-z .&]*?)#?\s+({NUM})\s+({NUM})\s+({NUM})\s+({NUM})\s*$")
BULLETIN_NAMES = {"Netherlands": "Netherland", "UK": "U.K.", "UAE": "U.A.E.", "Malaysia": "Malaysia*", "South Africa": "South Africa*",
                  "South Korea": "South Korea*", "Other Countries": "Other Countries#", "Other GCC Countries": "Other GCC Countries",
                  "EU Countries": "EU Countries"}


def read_bulletin() -> tuple[pd.DataFrame, dict]:
    """Annual remittances by country from SBP's bulletin: FY25, FY26, and July of FY26 and FY27 (provisional)."""
    pdf = pdfium.PdfDocument(str(config.RAW / config.BULLETIN_FILE))
    text = pdf[0].get_textpage().get_text_range()
    lines = [ln.strip() for ln in text.splitlines()]
    merged: list[str] = []
    for i, ln in enumerate(lines):                      # "15.Other Countries#" has its numbers on the next line
        if re.fullmatch(r"\d+\.Other Countries#", ln) and i + 1 < len(lines):
            ln = f"{ln} {lines[i + 1]}"
        merged.append(ln)
    values: dict[str, list[float]] = {}
    for ln in merged:
        m = ROW.match(ln)
        if m:
            name = m.group(1).strip()
            values[BULLETIN_NAMES.get(name, name)] = [float(g.replace(",", "")) for g in m.groups()[1:]]
    # the Total row reads "Total 38,299.2 ..." with no leading number
    needed = [*config.LEAF_TO_CORRIDOR, "Total", "U.A.E.", "EU Countries", "Other GCC Countries"]
    missing = [n for n in needed if n not in values]
    if missing:
        raise ValueError(f"the bulletin table lacks {missing}")
    cols = ["fy25", "fy26", "jul_fy26", "jul_fy27p"]
    wide = pd.DataFrame(values, index=cols).T

    checks = {}
    for col in cols:
        total = wide.loc["Total", col]
        for group, parts in (("U.A.E.", UAE_LEAVES), ("Other GCC Countries", OTHER_GCC_LEAVES), ("EU Countries", EU_LEAVES)):
            gap = abs(wide.loc[parts, col].sum() - wide.loc[group, col])
            if gap > 0.6:
                raise ValueError(f"bulletin {col}: {group} differs from its countries by {gap:.1f}")
        gap = abs(wide.loc[list(config.LEAF_TO_CORRIDOR), col].sum() - total)
        checks[f"max_gap_{col}"] = round(float(gap), 2)
        if gap > 1.0:
            raise ValueError(f"bulletin {col}: countries add to {wide.loc[list(config.LEAF_TO_CORRIDOR), col].sum():.1f}, total {total:.1f}")
    corridor = wide.loc[list(config.LEAF_TO_CORRIDOR)].groupby(config.LEAF_TO_CORRIDOR).sum().loc[config.CORRIDORS]
    corridor.index.name = "corridor"
    corridor = corridor.reset_index()
    corridor["region"] = corridor["corridor"].map(config.CORRIDOR_REGION)
    return corridor, {"bulletin_total": {c: float(wide.loc["Total", c]) for c in cols}, **checks}


def run() -> None:
    config.INTERIM.mkdir(parents=True, exist_ok=True)
    wide, fy_printed, checks = read_workbook()
    corridors = corridor_table(wide)
    bulletin, bulletin_checks = read_bulletin()
    wide.round(4).reset_index().to_csv(config.INTERIM / "monthly_by_sbp_row.csv", index=False, lineterminator="\n")
    corridors.round(4).to_csv(config.INTERIM / "monthly_by_corridor.csv", index=False, lineterminator="\n")
    bulletin.round(4).to_csv(config.INTERIM / "annual_by_corridor_bulletin.csv", index=False, lineterminator="\n")
    quality = {"workbook": checks, "bulletin": bulletin_checks, "series_break": "Country figures use the remitter's original country from July 2019."}
    (config.INTERIM / "data_quality.json").write_text(json.dumps(quality, indent=1) + "\n", encoding="utf-8", newline="\n")
    print(f"[build] {len(wide)} months, {len(config.CORRIDORS)} corridors; bulletin FY26 total {bulletin['fy26'].sum():,.1f} million US dollars")
